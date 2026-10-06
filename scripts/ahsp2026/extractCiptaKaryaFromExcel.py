import openpyxl
import re
import sys
import io
import json
import os

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

EXCEL_PATH = r"D:\file kerja\PEMBUATAN SOFTWARE\ezrab folder\ahsp exce\ahsp bina kontruksi 2026.xlsx"

print("Loading workbook from:", EXCEL_PATH)
wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
print("Workbook loaded successfully. Sheets:", len(wb.sheetnames))

# -------------------------------------------------------------
# 1. Extract Upah Bahan
# -------------------------------------------------------------
print("Extracting Upah Bahan...")
ws_upah_bahan = wb['Upah Bahan']
rows_ub = list(ws_upah_bahan.iter_rows(values_only=True))

ub_labor = []
ub_materials = []
ub_equipment = []

curr_ub_section = None
for r_idx, r in enumerate(rows_ub):
    col3 = str(r[3]).strip() if len(r) > 3 and r[3] is not None else ""
    col5 = str(r[5]).strip() if len(r) > 5 and r[5] is not None else ""
    
    if col3 == "I." or "UPAH" in col5:
        curr_ub_section = "labor"
        continue
    elif col3 == "II." or "MATERIAL" in col5:
        curr_ub_section = "materials"
        continue
    elif col3 in ["III", "III."] or "SEWA PERALATAN" in col5 or "PERALATAN" in col5:
        curr_ub_section = "equipment"
        continue
        
    kode = str(r[4] or "").strip() if len(r) > 4 else ""
    name = str(r[5] or "").strip() if len(r) > 5 else ""
    unit = str(r[6] or "").strip() if len(r) > 6 else ""
    price_val = r[7] if len(r) > 7 else None
    
    if not name or price_val is None or not isinstance(price_val, (int, float)) or price_val <= 0:
        continue
    if name.upper() in ["UPAH", "MATERIAL", "SEWA PERALATAN", "UPAH - MATERIAL - ALAT"]:
        continue
        
    entry = {
        "code": kode,
        "name": name,
        "unit": unit,
        "price": float(price_val)
    }
    
    if curr_ub_section == "labor":
        ub_labor.append(entry)
    elif curr_ub_section == "materials":
        ub_materials.append(entry)
    elif curr_ub_section == "equipment":
        ub_equipment.append(entry)

print(f"Upah Bahan extracted: labor={len(ub_labor)}, materials={len(ub_materials)}, equipment={len(ub_equipment)}")

# -------------------------------------------------------------
# 2. Extract DHSP
# -------------------------------------------------------------
print("Extracting Daftar Harga Satuan Pekerjaan (DHSP)...")
ws_dhsp = wb['Daftar Harga Satuan Pekerjaan']
dhsp_items = {}
for row in ws_dhsp.iter_rows(values_only=True):
    col1 = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
    norm = col1.rstrip(".")
    if re.match(r"^\d+(\.\d+)+[a-zA-Z]?$", norm) or re.match(r"^[A-Z]\.\d+(\.\d+)*[a-zA-Z]?$", norm):
        name = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
        unit = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ""
        price = row[4] if len(row) > 4 and isinstance(row[4], (int, float)) else 0
        dhsp_items[norm] = {
            "code": norm,
            "raw_code": col1,
            "name": name,
            "unit": unit,
            "price": float(price)
        }

print(f"DHSP extracted: {len(dhsp_items)} items")

# -------------------------------------------------------------
# 3. Extract Analysis Sheets (40 sheets)
# -------------------------------------------------------------
print("Extracting Analysis Sheets...")
analysis_sheets = [s for s in wb.sheetnames if s not in ["Daftar Harga Satuan Pekerjaan", "Upah Bahan"]]
parsed_analyses = {}

for sname in analysis_sheets:
    ws = wb[sname]
    rows = list(ws.iter_rows(values_only=True))
    r_idx = 0
    while r_idx < len(rows):
        r = rows[r_idx]
        c2 = str(r[2]).strip() if len(r) > 2 and r[2] is not None else ""
        norm_code = c2.rstrip(".")
        
        # Check if row looks like an item code
        if re.match(r"^\d+(\.\d+)+[a-zA-Z]?$", norm_code) or re.match(r"^[A-Z]\.\d+(\.\d+)*[a-zA-Z]?$", norm_code):
            # Look ahead up to 7 rows for "No" header row
            header_offset = None
            for la in range(1, 8):
                if r_idx + la < len(rows):
                    next_c2 = str(rows[r_idx + la][2]).strip() if len(rows[r_idx + la]) > 2 and rows[r_idx + la][2] is not None else ""
                    if next_c2.lower() == "no":
                        header_offset = la
                        break
            
            if header_offset is not None:
                desc = str(r[3]).strip() if len(r) > 3 and r[3] is not None else ""
                header_row = rows[r_idx + header_offset]
                col_map = {}
                for c_i, c_val in enumerate(header_row):
                    if c_val is None: continue
                    c_str = str(c_val).strip().lower()
                    if "uraian" in c_str: col_map["name"] = c_i
                    elif "kode" in c_str: col_map["code"] = c_i
                    elif "jumlah" in c_str: col_map["total"] = c_i
                    elif "harga" in c_str: col_map["price"] = c_i
                    elif "koef" in c_str: col_map["coeff"] = c_i
                    elif "sat" in c_str: col_map["unit"] = c_i
                
                cur_r = r_idx + header_offset + 1
                item_data = {
                    "code": norm_code,
                    "raw_code": c2,
                    "description": desc,
                    "unit": dhsp_items.get(norm_code, {}).get("unit", ""),
                    "unit_price": dhsp_items.get(norm_code, {}).get("price", 0.0),
                    "labor": [],
                    "materials": [],
                    "equipment": [],
                    "sheet": sname
                }
                
                curr_cat = None
                while cur_r < len(rows):
                    row_data = rows[cur_r]
                    rc2 = str(row_data[2]).strip() if len(row_data) > 2 and row_data[2] is not None else ""
                    rc3 = str(row_data[3]).strip() if len(row_data) > 3 and row_data[3] is not None else ""
                    
                    if rc2 == "A" or "tenaga" in rc3.lower():
                        curr_cat = "labor"
                        cur_r += 1
                        continue
                    elif rc2 == "B" or "bahan" in rc3.lower():
                        curr_cat = "materials"
                        cur_r += 1
                        continue
                    elif rc2 == "C" or "peralatan" in rc3.lower() or "alat" in rc3.lower():
                        curr_cat = "equipment"
                        cur_r += 1
                        continue
                    elif rc2 in ["D", "E", "F"] or "jumlah (a+b+c)" in rc3.lower() or "biaya umum" in rc3.lower() or "harga satuan pekerjaan" in rc3.lower():
                        if rc2 == "F" or "harga satuan pekerjaan" in rc3.lower():
                            if not item_data["unit_price"]:
                                for val in row_data[4:]:
                                    if isinstance(val, (int, float)) and val > 0:
                                        item_data["unit_price"] = float(val)
                                        break
                            cur_r += 1
                            break
                        cur_r += 1
                        continue
                        
                    if not rc2 and not rc3:
                        if cur_r + 2 < len(rows) and str(rows[cur_r + 2][2]).strip().lower() == "no":
                            cur_r += 1
                            break
                        cur_r += 1
                        continue
                        
                    # Check if next item header is reached
                    if cur_r + 1 < len(rows):
                        next_rc2 = str(rows[cur_r + 1][2]).strip() if len(rows[cur_r + 1]) > 2 and rows[cur_r + 1][2] is not None else ""
                        if (re.match(r"^\d+(\.\d+)+[a-zA-Z]?$", next_rc2.rstrip(".")) or re.match(r"^[A-Z]\.\d+(\.\d+)*[a-zA-Z]?$", next_rc2.rstrip("."))) and cur_r + 2 < len(rows) and str(rows[cur_r + 2][2]).strip().lower() == "no":
                            break
                        
                    if curr_cat and (re.match(r"^\d+$", rc2) or (rc3 and not rc2.startswith("JUMLAH"))):
                        if "JUMLAH HARGA" in rc2 or "JUMLAH HARGA" in rc3:
                            cur_r += 1
                            continue
                        comp_name = str(row_data[col_map.get("name", 3)] or "").strip()
                        comp_code = str(row_data[col_map["code"]] or "").strip() if "code" in col_map and len(row_data) > col_map["code"] else ""
                        comp_unit = str(row_data[col_map["unit"]] or "").strip() if "unit" in col_map and len(row_data) > col_map["unit"] else ""
                        comp_coeff = row_data[col_map["coeff"]] if "coeff" in col_map and len(row_data) > col_map["coeff"] else 0
                        comp_price = row_data[col_map["price"]] if "price" in col_map and len(row_data) > col_map["price"] else 0
                        comp_total = row_data[col_map["total"]] if "total" in col_map and len(row_data) > col_map["total"] else 0
                        
                        if comp_name and comp_name != "Uraian":
                            item_data[curr_cat].append({
                                "code": comp_code,
                                "name": comp_name,
                                "unit": comp_unit,
                                "coefficient": float(comp_coeff) if isinstance(comp_coeff, (int, float)) else 0.0,
                                "unitPrice": float(comp_price) if isinstance(comp_price, (int, float)) else 0.0,
                                "total": float(comp_total) if isinstance(comp_total, (int, float)) else 0.0
                            })
                    cur_r += 1
                
                parsed_analyses[norm_code] = item_data
                r_idx = cur_r
                continue
        r_idx += 1

print(f"Total analysis items parsed: {len(parsed_analyses)}")

out_dir = os.path.join(os.getcwd(), "data", "ahsp2026", "excel_extracted")
os.makedirs(out_dir, exist_ok=True)

with open(os.path.join(out_dir, "cipta_karya_analyses.json"), "w", encoding="utf-8") as f:
    json.dump(parsed_analyses, f, indent=2, ensure_ascii=False)

with open(os.path.join(out_dir, "dhsp_items.json"), "w", encoding="utf-8") as f:
    json.dump(dhsp_items, f, indent=2, ensure_ascii=False)

with open(os.path.join(out_dir, "upah_bahan.json"), "w", encoding="utf-8") as f:
    json.dump({
        "labor": ub_labor,
        "materials": ub_materials,
        "equipment": ub_equipment
    }, f, indent=2, ensure_ascii=False)

print("Saved extracted files to:", out_dir)
