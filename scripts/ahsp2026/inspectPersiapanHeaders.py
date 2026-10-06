import openpyxl
import re

wb = openpyxl.load_workbook('D:/file kerja/PEMBUATAN SOFTWARE/ezrab folder/ahsp exce/ahsp bina kontruksi 2026.xlsx', data_only=True)

# Check sheet Persiapan
ws = wb['Persiapan']
print("Scanning Persiapan for item headers:")
for r in range(1, 100):
    val_c = ws.cell(r, 3).value # Column C (index 3)
    val_d = ws.cell(r, 4).value # Column D (index 4)
    if val_c is not None:
        s_c = str(val_c).strip()
        # If it looks like a number/code
        if re.match(r'^\d+(\.\d+)*$', s_c):
            # Check what's in the next row
            next_c = ws.cell(r+1, 3).value
            print(f"Row {r:3d}: Col C='{s_c:12s}' | Col D='{str(val_d)[:50]}' | Next Row Col C='{next_c}'")
