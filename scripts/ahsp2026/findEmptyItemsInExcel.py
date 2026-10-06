import openpyxl
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

EXCEL_PATH = r"D:\file kerja\PEMBUATAN SOFTWARE\ezrab folder\ahsp exce\ahsp bina kontruksi 2026.xlsx"
wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)

search_terms = ["4.1.1.1", "4.1.1.2", "8.1.8.1", "8.3.8.3", "9.4.1.1", "U-Ditch 30x40", "Box Culvert 80x80", "Polybag 20 Liter", "Pipa DCI, DN. 5"]

for sname in wb.sheetnames:
    ws = wb[sname]
    for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
        row_str = " | ".join(str(c) for c in row if c is not None)
        for term in search_terms:
            if term.lower() in row_str.lower():
                print(f"Sheet '{sname}' Row {r_idx+1}: matched '{term}' -> {row_str[:120]}")
