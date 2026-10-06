import openpyxl

wb = openpyxl.load_workbook(r'D:\file kerja\PEMBUATAN SOFTWARE\ezrab folder\ahsp exce\AHSP 2026 Bina Marga.xlsx', data_only=False)

items_to_audit = [
    ('2.1.(1)', 'B - Drainase', 2, 'Drainase'),
    ('2.3.(1)', 'B - Drainase', 189, 'Drainase'),
    ('3.1.(1)', 'C - Tanah dan Geosintetik', 2, 'Tanah'),
    ('5.1.(1a)', 'E - Perkerasan Berbutir dan Per', 2, 'Perkerasan Berbutir'),
    ('6.1.(1)', 'F - Perkerasan Aspal', 2, 'Aspal'),
    ('7.1.(1a2)', 'G - Struktur', 2, 'Struktur'),
]

for code, sheet_name, start_row, domain_cat in items_to_audit:
    ws = wb[sheet_name]
    print('=' * 80)
    print(f'FORENSIC TRACE FOR {code} ({domain_cat}) IN SHEET {sheet_name}')
    print('=' * 80)
    
    header_name = ws.cell(start_row, 4).value
    header_unit = ws.cell(start_row, 6).value
    header_dhsp_formula = ws.cell(start_row, 11).value
    print(f'AHSP CODE: {code}')
    print(f'AHSP NAME: {header_name}')
    print(f'UNIT: {header_unit}')
    print(f'DHSP FORMULA: {header_dhsp_formula}')
    print('-' * 80)
    print(f'{"TYPE":<10} | {"CODE":<8} | {"RESOURCE NAME":<35} | {"COEF":<10} | {"UNIT":<6} | {"PRICE FORMULA / VAL":<30}')
    print('-' * 80)
    
    curr_type = 'UNKNOWN'
    r = start_row + 1
    while r < start_row + 50:
        c_k = ws.cell(r, 11).value
        if r > start_row + 3 and c_k and str(c_k).startswith('=ROUNDDOWN'):
            break
            
        c_c = str(ws.cell(r, 3).value or '').strip()
        c_d = str(ws.cell(r, 4).value or '').strip()
        c_e = str(ws.cell(r, 5).value or '').strip()
        c_f = str(ws.cell(r, 6).value or '').strip()
        c_g = ws.cell(r, 7).value # Coef
        c_h = ws.cell(r, 8).value # Unit price
        c_j = ws.cell(r, 10).value # Subtotal
        
        if c_c == 'A' and 'TENAGA' in c_d.upper():
            curr_type = 'LABOR'
        elif c_c == 'B' and 'BAHAN' in c_d.upper():
            curr_type = 'MATERIAL'
        elif c_c == 'C' and 'PERALATAN' in c_d.upper():
            curr_type = 'EQUIPMENT'
        elif c_c == 'D' and 'JUMLAH' in c_d.upper():
            print(f'{"DIRECT COST (D)":<40} : formula={c_j}')
        elif c_c == 'E' and 'BIAYA UMUM' in c_d.upper():
            print(f'{"OVERHEAD (E)":<40} : formula={c_j}')
        elif c_c == 'F' and 'HARGA SATUAN' in c_d.upper():
            print(f'{"HSP (F)":<40} : formula={c_j}')
        elif c_g is not None and c_g != '':
            print(f'{curr_type:<10} | {c_e:<8} | {c_d[:35]:<35} | {str(c_g)[:10]:<10} | {c_f:<6} | {str(c_h)[:30]:<30}')
            
        r += 1
    print('\n')
