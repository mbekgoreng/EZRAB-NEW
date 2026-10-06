import openpyxl

wb = openpyxl.load_workbook(r'D:\file kerja\PEMBUATAN SOFTWARE\ezrab folder\ahsp exce\AHSP 2026 Bina Marga.xlsx', data_only=False)
ws_b = wb['B - Drainase']

for c in range(1, 15):
    val = ws_b.cell(2, c).value
    col_letter = openpyxl.utils.get_column_letter(c)
    print(f'Col {col_letter} (c={c}): {repr(val)}')
