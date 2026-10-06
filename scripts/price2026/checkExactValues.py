import openpyxl

wb_val = openpyxl.load_workbook(r'D:\file kerja\PEMBUATAN SOFTWARE\ezrab folder\ahsp exce\AHSP 2026 Bina Marga.xlsx', data_only=True)
ws_b = wb_val['B - Drainase']

# Check row 15 (Direct cost), row 16 (Overhead), row 17 (HSP), row 2 (DHSP)
print('Row 7 (Labor sum):', ws_b.cell(7, 10).value)
print('Row 14 (Equip sum):', ws_b.cell(14, 10).value)
print('Row 15 (Direct cost D):', ws_b.cell(15, 10).value)
print('Row 16 (Overhead E):', ws_b.cell(16, 10).value)
print('Row 17 (HSP F):', ws_b.cell(17, 10).value)
print('Row 2 col K (DHSP):', ws_b.cell(2, 11).value)
