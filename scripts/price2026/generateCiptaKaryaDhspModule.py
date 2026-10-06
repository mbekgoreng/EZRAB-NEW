import openpyxl
import re
import os
import json

WORKBOOK_PATH = r"D:\file kerja\PEMBUATAN SOFTWARE\ezrab folder\ahsp exce\ahsp bina kontruksi 2026.xlsx"
OUT_TS = os.path.join("src", "data", "nationalCostDatabase", "officialCiptaKaryaDhsp2026.ts")

print("Loading workbook for DHSP generation...")
wb_values = openpyxl.load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
wb_formulas = openpyxl.load_workbook(WORKBOOK_PATH, read_only=True, data_only=False)

ws_v = wb_values['Daftar Harga Satuan Pekerjaan']
ws_f = wb_formulas['Daftar Harga Satuan Pekerjaan']

rows_v = list(ws_v.iter_rows(values_only=True))
rows_f = list(ws_f.iter_rows(values_only=True))

dhsp_records = {}
leaf_count = 0
header_count = 0

for r_idx in range(4, len(rows_v)):
    r_v = rows_v[r_idx]
    r_f = rows_f[r_idx]
    
    col_code = str(r_v[1]).strip() if len(r_v) > 1 and r_v[1] is not None else ""
    col_name = str(r_v[2]).strip() if len(r_v) > 2 and r_v[2] is not None else ""
    col_unit = str(r_v[3]).strip() if len(r_v) > 3 and r_v[3] is not None else ""
    price_val = r_v[4] if len(r_v) > 4 else None
    ket = str(r_v[5]).strip() if len(r_v) > 5 and r_v[5] is not None else ""
    
    formula_raw = str(r_f[4]) if len(r_f) > 4 and r_f[4] is not None else ""
    
    norm = col_code.rstrip(".")
    if not norm:
        continue
        
    if re.match(r"^\d+(\.\d+)+[a-zA-Z]?$", norm) or re.match(r"^[A-Z]\.\d+(\.\d+)*[a-zA-Z]?$", norm):
        price_num = float(price_val) if isinstance(price_val, (int, float)) and price_val > 0 else None
        
        is_leaf = bool(col_unit and price_num is not None)
        if is_leaf:
            leaf_count += 1
        else:
            header_count += 1
            
        dhsp_records[norm] = {
            "code": norm,
            "rawCode": col_code,
            "name": col_name,
            "unit": col_unit,
            "unitPrice": price_num,
            "isLeaf": is_leaf,
            "sourceSheet": "Daftar Harga Satuan Pekerjaan",
            "sourceRow": r_idx + 1,
            "formulaRaw": formula_raw if formula_raw.startswith("=") else None,
            "notes": ket or None
        }

print(f"Extracted {len(dhsp_records)} DHSP entries (Leaf: {leaf_count}, Headers: {header_count})")

# Write TypeScript module
lines = []
lines.append("/**")
lines.append(" * OFFICIAL CIPTA KARYA DHSP 2026 (DAFTAR HARGA SATUAN PEKERJAAN)")
lines.append(" * Sourced directly from SE DJBK No. 47/SE/Dk/2026 (Workbook: ahsp bina kontruksi 2026.xlsx)")
lines.append(" * Evaluated from sheet 'Daftar Harga Satuan Pekerjaan' with formula traceability")
lines.append(" */")
lines.append("")
lines.append("export interface OfficialCiptaKaryaDhspEntry {")
lines.append("  code: string;")
lines.append("  rawCode: string;")
lines.append("  name: string;")
lines.append("  unit: string;")
lines.append("  unitPrice: number | null;")
lines.append("  isLeaf: boolean;")
lines.append("  sourceSheet: string;")
lines.append("  sourceRow: number;")
lines.append("  formulaRaw?: string | null;")
lines.append("  notes?: string | null;")
lines.append("}")
lines.append("")

# Chunk to avoid TS compiler recursion
chunk_size = 250
items_list = list(dhsp_records.values())
chunks = []
for i in range(0, len(items_list), chunk_size):
    cname = f"__DHSP_CHUNK_{i // chunk_size}"
    s_json = json.dumps(items_list[i:i+chunk_size], indent=2, ensure_ascii=False)
    lines.append(f"const {cname}: OfficialCiptaKaryaDhspEntry[] = {s_json};")
    chunks.append(cname)
    lines.append("")

lines.append("export const OFFICIAL_CK_2026_DHSP_LIST: OfficialCiptaKaryaDhspEntry[] = [")
for cn in chunks:
    lines.append(f"  ...{cn},")
lines.append("];")
lines.append("")
lines.append("export const OFFICIAL_CK_2026_DHSP_MAP = new Map<string, OfficialCiptaKaryaDhspEntry>(")
lines.append("  OFFICIAL_CK_2026_DHSP_LIST.map((entry) => [entry.code, entry])")
lines.append(");")
lines.append("")
lines.append(f"export const CK_2026_DHSP_TOTAL_COUNT = {len(dhsp_records)};")
lines.append(f"export const CK_2026_DHSP_LEAF_COUNT = {leaf_count};")
lines.append("")

with open(OUT_TS, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))

print(f"Generated {OUT_TS} with {len(dhsp_records)} entries!")
